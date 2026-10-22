#!/usr/bin/env python3
"""把仓库根目录的 README.md 转成 DOCX，方便发给不看 GitHub 的人。

这是手动用的小工具，不参与自动化流程，也不在 requirements.txt 里：
先自行安装依赖 `pip install python-docx`，再运行

    .venv/bin/python scripts/convert_readme_to_docx.py [输出路径]

不传输出路径时写到系统临时目录（$TMPDIR 或 /tmp），避免在仓库里留下产物。
"""

import os
import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE


def parse_markdown_table(lines):
    """Parse markdown table into list of rows."""
    rows = []
    for line in lines:
        line = line.strip()
        if not line or line.startswith('|---') or line.startswith('| ---'):
            continue
        if line.startswith('|') and line.endswith('|'):
            cells = [cell.strip() for cell in line.split('|')[1:-1]]
            rows.append(cells)
    return rows


def add_table_to_doc(doc, rows):
    """Add table to document."""
    if not rows:
        return

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = 'Table Grid'

    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = cell_text

            # Bold header row
            if i == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def convert_readme_to_docx(readme_path, output_path):
    """Convert README.md to DOCX."""
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(10)

    # Read README
    with open(readme_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_lines = []
    table_lines = []

    while i < len(lines):
        line = lines[i]

        # Handle code blocks
        if line.startswith('```'):
            if in_code_block:
                # End code block
                code_text = '\n'.join(code_lines)
                p = doc.add_paragraph()
                p.style = doc.styles['Normal']
                run = p.add_run(code_text)
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
                code_lines = []
                in_code_block = False
            else:
                # Start code block
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle tables
        if line.strip().startswith('|') and line.strip().endswith('|'):
            table_lines.append(line)
            i += 1
            # Check if next line is still part of table
            if i < len(lines) and lines[i].strip().startswith('|'):
                continue
            else:
                # End of table
                if table_lines:
                    rows = parse_markdown_table(table_lines)
                    add_table_to_doc(doc, rows)
                    table_lines = []
                continue

        # Handle headings
        if line.startswith('#'):
            level = len(line.split(' ')[0])
            text = line.lstrip('#').strip()

            if level == 1:
                heading = doc.add_heading(text, level=0)
            elif level == 2:
                heading = doc.add_heading(text, level=1)
            elif level == 3:
                heading = doc.add_heading(text, level=2)
            else:
                heading = doc.add_heading(text, level=3)

            i += 1
            continue

        # Handle bullet points
        if line.strip().startswith('- ') or line.strip().startswith('* '):
            text = line.strip()[2:]
            # Clean markdown formatting
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            text = re.sub(r'`(.*?)`', r'\1', text)

            p = doc.add_paragraph(text, style='List Bullet')
            i += 1
            continue

        # Handle numbered lists
        if re.match(r'^\d+\.', line.strip()):
            text = re.sub(r'^\d+\.\s*', '', line.strip())
            # Clean markdown formatting
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            text = re.sub(r'`(.*?)`', r'\1', text)

            p = doc.add_paragraph(text, style='List Number')
            i += 1
            continue

        # Handle empty lines
        if not line.strip():
            i += 1
            continue

        # Handle regular paragraphs
        text = line.strip()
        # Clean markdown formatting
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
        text = re.sub(r'`(.*?)`', r'\1', text)
        text = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', text)

        if text:
            p = doc.add_paragraph(text)

        i += 1

    # Save document
    doc.save(output_path)
    print(f"Converted {readme_path} to {output_path}")


if __name__ == '__main__':
    readme_path = Path(__file__).resolve().parent.parent / 'README.md'
    if len(sys.argv) > 1:
        output_path = Path(sys.argv[1]).expanduser()
    else:
        temp_dir = Path(os.environ.get('TMPDIR') or '/tmp')
        output_path = temp_dir / '115-media-hub-README.docx'

    if readme_path.exists():
        convert_readme_to_docx(readme_path, output_path)
    else:
        print(f"README.md not found at {readme_path}")
